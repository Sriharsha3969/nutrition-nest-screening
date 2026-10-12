import re
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

COLOR_HIERARCHY = {
    "Red": {"priority": 1, "status": "Rejected", "fill": "FF7C80", "font": "721C24", "prefix": "Banned ingredient"},
    "Purple": {"priority": 2, "status": "Rejected", "fill": "CDB4F6", "font": "3B1E6D", "prefix": "Undisclosed formula"},
    "Orange": {"priority": 3, "status": "Rejected", "fill": "FFC000", "font": "7F6000", "prefix": "Not US-distributed"},
    "Blue": {"priority": 4, "status": "Rejected", "fill": "9DC3E6", "font": "1F3864", "prefix": "Premium fit / rating below 4.2"},
    "Pink": {"priority": 5, "status": "Rejected", "fill": "F4B6C2", "font": "721C24", "prefix": "Unsubstantiated health claim"},
    "Grey": {"priority": 6, "status": "Rejected", "fill": "BFBFBF", "font": "333333", "prefix": "Duplicate listing"},
    "Yellow": {"priority": 7, "status": "Pending Data", "fill": "FFFF99", "font": "7F6000", "prefix": "Pending data"},
    "Amber": {"priority": 8, "status": "Hold (Proof Required)", "fill": "FFE699", "font": "856404", "prefix": "Hold - proof required"},
    "None": {"priority": 9, "status": "Passed", "fill": "C6EFCE", "font": "006100", "prefix": "Passed - no flags"}
}

G01_PLACEHOLDER_PATTERNS = [
    r"^\s*$", r"^nan$", r"^n/?a$", r"^none$", r"^null$", r"^customize$",
    r"verify on label", r"varies by", r"variety pack", r"variety bundle",
    r"assorted.*flavors", r"not (?:successfully )?retrieved",
    r"labels? not retrieved", r"not exposed", r"not found in sources",
    r"not currently published", r"does not display a formal",
    r"complete (?:component |individual )?ingredient.*not ",
    r"exact label wording", r"\bassumed\b", r"\bpresumed\b", r"same recipe as"
]

BANNED_PATTERNS = [
    ("F-01", r"\b(?:potassium bromate|bromated flour|e924)\b", "potassium bromate ban"),
    ("F-02", r"\b(?:azodicarbonamide|e927a|dough conditioner)\b", "azodicarbonamide ban"),
    ("F-03", r"\b(?:brominated vegetable oil|bvo|brominated oil)\b", "BVO ban"),
    ("F-04", r"\b(?:fd&c|red\s*(?:3|40)|yellow\s*(?:5|6)|blue\s*(?:1|2)|green\s*3|e102|e110|e122|e124|e129|e133|erythrosine|tartrazine|allura red|brilliant blue|sunset yellow)\b", "synthetic petroleum-dye ban"),
    ("F-05", r"\b(?:titanium dioxide|tio2|e171|ci\s*77891)\b", "titanium dioxide ban"),
    ("F-06", r"\b(?:sucralose|aspartame|acesulfame(?:-k|\s*potassium)?|saccharin|sodium saccharin|cyclamate|neotame|advantame|e950|e951|e952|e954|e955)\b", "artificial sweetener ban"),
    ("F-07", r"\b(?:partially hydrogenated|interesterified fat|hydrogenated (?:soybean|cottonseed|palm|canola|vegetable) oil|shortening)\b", "trans fat / hydrogenated oil ban"),
    ("F-08", r"\b(?:sodium nitrite|sodium nitrate|potassium nitrite|potassium nitrate|e249|e250|e251|e252)\b", "nitrite / nitrate preservative ban"),
    ("F-09", r"\b(?:bha|bht|tbhq|propyl gallate|ethoxyquin|butylated hydroxyanisole|butylated hydroxytoluene|tert-butylhydroquinone|e310|e319|e320|e321)\b", "synthetic antioxidant preservative ban"),
    ("F-11", r"\b(?:caramel colou?r|e150c|e150d|class (?:iii|iv) caramel)\b", "caramel color ban"),
    ("F-12", r"\b(?:propylparaben|methylparaben|e216|e217)\b", "paraben preservative ban"),
    ("F-13", r"\b(?:high[- ]fructose corn syrup|hfcs|glucose-fructose syrup|corn syrup(?:\s*solids)?)\b", "HFCS / refined sweetener syrup ban"),
    ("F-15", r"\b(?:methylene chloride|dichloromethane|solvent-decaffeinated)\b", "methylene-chloride decaf ban"),
    ("F-18", r"\b(?:swordfish|king mackerel|shark|tilefish|bigeye tuna|marlin)\b", "high-mercury fish ban"),
    ("F-19", r"\b(?:dmaa|dmba|dmha|sibutramine|ephedra|ephedrine|sarms|prohormone|phenibut|tianeptine|kratom|aristolochic acid|comfrey)\b", "banned supplement compound"),
    ("F-20", r"\b(?:talc|magnesium silicate)\b", "talc / asbestos risk ban"),
    ("F-21", r"\b(?:soybean oil|canola oil|rapeseed oil|cottonseed oil|corn oil|hexane-extracted)\b", "refined industrial seed oil"),
    ("P-01", r"\b(?:methylparaben|ethylparaben|propylparaben|butylparaben|isobutylparaben|benzylparaben)\b", "paraben endocrine disruptor ban"),
    ("P-02", r"\b(?:dibutyl phthalate|diethylhexyl phthalate|dbp|dehp|dep|bbp|phthalate)\b", "phthalate ban"),
    ("P-03", r"\b(?:formaldehyde|dmdm hydantoin|quaternium-15|imidazolidinyl urea|diazolidinyl urea|bronopol|sodium hydroxymethylglycinate)\b", "formaldehyde / releaser ban"),
    ("P-04", r"\b(?:ptfe|teflon|perfluoro\w+|polyfluoro\w+)\b", "PFAS / forever-chemical ban"),
    ("P-05", r"\b(?:triclosan|triclocarban|chloroxylenol)\b", "triclosan antibacterial ban"),
    ("P-07", r"\b(?:hydroquinone|lead acetate|coal tar|thimerosal)\b", "hydroquinone / coal-tar ban"),
    ("P-08", r"\b(?:p-phenylenediamine|ppd|toluene-2,5-diamine|methylene glycol)\b", "aromatic amine / hair straightener ban"),
    ("P-09", r"\b(?:oxybenzone|benzophenone-3|octinoxate|homosalate|octocrylene)\b", "endocrine-disrupting chemical UV filter ban"),
    ("P-10", r"\b(?:benzoyl peroxide)\b", "benzene-forming peroxide ban"),
    ("P-11", r"\b(?:peg-\d+|ppg-\d+|sodium laureth sulfate|sles|ceteareth-\d+|polysorbate[- ]\d+|poloxamer[- ]\d+)\b", "ethoxylated ingredient (1,4-dioxane risk)"),
    ("P-12", r"\b(?:cocamide dea|lauramide dea|triethanolamine|diethanolamine)\b", "nitrosamine-forming amine ban"),
    ("P-13", r"\b(?:cyclotetrasiloxane|cyclopentasiloxane|cyclohexasiloxane)\b", "cyclic siloxane (D4/D5/D6) ban"),
    ("P-14", r"\b(?:nylon-12|acrylates copolymer|pet glitter|polyethylene beads)\b", "microplastic ban"),
    ("P-15", r"\b(?:sodium lauryl sulfate|sls|methylisothiazolinone|methylchloroisothiazolinone)\b", "harsh sulfate (SLS) / MIT sensitizer ban"),
    ("P-16", r"\b(?:aluminum chlorohydrate|aluminum zirconium|galaxolide|tonalide|musk ketone)\b", "aluminum salt / polycyclic musk ban"),
    ("H-01", r"\b(?:paraffin wax|petroleum paraffin)\b", "paraffin candle wax ban"),
    ("H-04", r"\b(?:benzalkonium chloride|didecyldimethylammonium chloride)\b", "quaternary ammonium (quat) ban"),
    ("H-05", r"\b(?:sodium hypochlorite|chlorine bleach|2-butoxyethanol|perchloroethylene)\b", "chlorine bleach / glycol ether ban"),
    ("H-06", r"\b(?:nonylphenol|optical brightener)\b", "NPE / optical brightener ban")
]

F14_MARKERS = [
    ("mono-/diglycerides", r"\bmono-?\s*(?:and|&|/)?\s*diglycerides(?:\s*of fatty acids)?\b"),
    ("polysorbate", r"\bpolysorbate\s*\d+\b"),
    ("carboxymethylcellulose", r"\b(?:carboxymethyl\s*cellulose|cellulose gum|cmc|e466)\b"),
    ("artificial flavor", r"\b(?:artificial\s+(?:\w+\s+)?flavo(?:u)?rs?(?:ing)?|nature-identical\s+flavo(?:u)?rs?)\b"),
    ("flavor enhancer (MSG/E621-E635)", r"\b(?:monosodium glutamate|msg|disodium inosinate|disodium guanylate|e621|e627|e631|e635)\b"),
    ("modified food starch", r"\bmodified\s+(?:corn|food|potato)\s+starch\b"),
    ("soy lecithin (conventional)", r"(?<!organic\s)\bsoy lecithin\b")
]

G07_CLAIM_PATTERNS = [
    r"\bcures?\s+(?:cancer|diabetes|alzheimer|disease|autism|arthritis)\b",
    r"\bmiracle\s+weight[- ]loss\b",
    r"\b100%\s+disease\s+reversal\b"
]

W01_SUGARS = [
    r"\bcane sugar\b", r"\braw cane sugar\b", r"\bhoney\b",
    r"\bagave(?:\s*nectar|\s*syrup)?\b", r"\bmaple syrup\b", r"\bdate syrup\b",
    r"\bcoconut sugar\b", r"\bbrown sugar\b", r"\b(?:\w+\s+)?juice concentrate\b",
    r"(?<!zero\s)(?<!no\s)(?<!without\s)\bsugar\b"
]

W02_CAFFEINE = [
    r"\bcaffeine\b", r"\bguarana\b", r"\byerba mate\b", r"\bcoffeeberry\b",
    r"\bsvetol\b", r"\bgreen coffee(?:\s*bean)?\b", r"\bkola nut\b", r"\bgreen tea caffeine\b"
]

W03_GUMS = [
    r"\bcarrageenan\b", r"\be407\b", r"\bxanthan(?:\s*gum)?\b",
    r"\bguar(?:\s*gum|\s*fiber)?\b", r"\bgellan(?:\s*gum)?\b",
    r"\barabic gum\b", r"\bgum acacia\b", r"\bacacia gum\b",
    r"\blocus bean gum\b", r"\btara gum\b"
]

G06_HEAVY_METAL_TRIGGERS = [
    ("tea", r"\b(?:green|black|white|oolong|herbal|mint|chamomile|echinacea)?\s*tea\b"),
    ("matcha", r"\bmatcha\b"),
    ("coffee", r"\bcoffee\b"),
    ("cacao", r"\bcacao\b"),
    ("cocoa", r"\bcocoa\b"),
    ("chocolate", r"\bchocolate\b"),
    ("turmeric", r"\bturmeric\b"),
    ("cinnamon", r"\bcinnamon\b"),
    ("pepper", r"\b(?:black|cayenne|chili|red|white)\s+pepper\b"),
    ("cardamom", r"\bcardamom\b"),
    ("clove", r"\bcloves?\b"),
    ("ginger", r"\bginger\b"),
    ("rice", r"\brice(?:\s*flour|\s*syrup|\s*starch|\s*protein)?\b"),
    ("seaweed", r"\b(?:seaweed|kelp|nori|wakame|kombu|dulse)\b"),
    ("sea moss", r"\bsea moss\b"),
    ("spirulina", r"\bspirulina\b"),
    ("chlorella", r"\bchlorella\b"),
    ("protein powder", r"\b(?:whey|pea|hemp|brown rice|soy)\s+protein\b")
]

F17_AUTHENTICITY_TRIGGERS = [
    "turmeric", "cinnamon", "honey", "chili", "paprika", "saffron", "extra virgin olive oil"
]

W06_CONVENTIONAL_CROPS = [
    "wheat", "oats", "oat", "lentils", "lentil", "chickpeas", "chickpea",
    "potatoes", "potato", "apple", "apples", "strawberry", "strawberries",
    "spinach", "kale", "grape", "grapes", "peach", "peaches", "pear", "pears",
    "nectarine", "cherry", "cherries", "blueberry", "blueberries",
    "bell pepper", "celery", "tomato", "tomatoes", "ginger", "matcha",
    "coffee", "sencha", "oolong", "black tea", "green tea", "cardamom"
]

W12_UPPER_LIMIT_NUTRIENTS = [
    ("vitamin a", r"\bvitamin a\b|\bretinyl palmitate\b"),
    ("vitamin b6", r"\bvitamin b6\b|\bpyridoxine\b"),
    ("zinc", r"\bzinc\b"),
    ("iron", r"\biron\b|\bferrous\b"),
    ("melatonin", r"\bmelatonin\b"),
    ("niacin", r"\bniacin\b(?!amide)")
]

W18_BOTANICAL_CAUTIONS = [
    ("ashwagandha", r"\bashwagandha\b|\bwithania\b"),
    ("green tea extract", r"\bgreen tea extract\b|\begcg\b"),
    ("licorice", r"\blicorice\b|\bglycyrrhiza\b"),
    ("black cohosh", r"\bblack cohosh\b"),
    ("kava", r"\bkava\b"),
    ("st. john's wort", r"\bst\.?\s*john'?s wort\b")
]


def detect_w06_conventional(ing_text: str, prod_name: str = "") -> list:
    if not ing_text:
        return []
    clauses = re.split(r"[,;()\[\]]|\band\b", ing_text, flags=re.I)
    if re.search(r"\b(?:100%\s+organic|usda\s+organic)\b", prod_name, re.I):
        return []
    found_crops = []
    for clause in clauses:
        c_low = clause.strip().lower()
        if not c_low or "organic" in c_low:
            continue
        for crop in W06_CONVENTIONAL_CROPS:
            if re.search(r"\b" + re.escape(crop.lower()) + r"\b", c_low):
                if not re.search(r"\borganic\b.*" + re.escape(crop.lower()), prod_name.lower()):
                    norm_crop = crop.lower().rstrip("s") if crop.lower() not in ["oats", "lentils", "potatoes", "molasses"] else crop.lower()
                    if norm_crop not in found_crops:
                        found_crops.append(norm_crop)
    return found_crops


def evaluate_product_comprehensive(
    row: dict,
    min_rating: float = 4.2,
    is_exact_dup: bool = False,
    dup_of_row: int = None,
    sibling_rows: list = None,
    prefer_existing_notes: bool = False
):
    existing_issue = str(row.get("Screening issue", "")).strip()
    if prefer_existing_notes and existing_issue and existing_issue.lower() not in ["nan", "none", ""]:
        low = existing_issue.lower()
        if low.startswith("banned"):
            return "Rejected", "Red", existing_issue
        if low.startswith("undisclosed"):
            return "Rejected", "Purple", existing_issue
        if low.startswith("not us-distributed"):
            return "Rejected", "Orange", existing_issue
        if low.startswith("premium fit"):
            return "Rejected", "Blue", existing_issue
        if low.startswith("pending data"):
            return "Pending Data", "Yellow", existing_issue
        if low.startswith("hold"):
            return "Hold (Proof Required)", "Amber", existing_issue
        if low.startswith("passed"):
            return "Passed", "None", existing_issue

    cat = str(row.get("Category", "")).strip()
    cat_l = cat.lower()
    p_name = str(row.get("Product name", "")).strip()
    ing_raw = row.get("Ingredients", None)
    ing = "" if pd.isna(ing_raw) else str(ing_raw).strip()
    ing_l = ing.lower()
    desc = "" if pd.isna(row.get("Product description")) else str(row.get("Product description")).strip()
    sizes = "" if pd.isna(row.get("Sizes available")) else str(row.get("Sizes available")).strip()
    mfg_link = "" if pd.isna(row.get("Manufacturer link")) else str(row.get("Manufacturer link")).strip()

    is_hardware = any(k in cat_l for k in ["wearables", "fitness & recovery", "travel"]) or bool(
        re.search(r"\b(?:toothbrush|brush head|squeezer|floss pick|scraper|diffuser device)\b", p_name, re.I)
    )
    is_ingestible = any(
        k in cat_l for k in ["grocery", "snack", "beverage", "drink", "coffee", "tea", "sport", "vitamin", "longevity", "sleep", "food", "pantry", "nutrition", "supplement", "protein"]
    ) or bool(
        re.search(r"\b(?:flour|sugar|salt|cocoa|milk|whey|syrup|starch|cheese|butter|oats|wheat)\b", ing_l)
    )
    is_supplement = any(k in cat_l for k in ["sport", "vitamin", "longevity", "sleep", "supplement"])
    is_oral = "oral" in cat_l

    triggered_colors = []
    hard_issues = []
    hold_issues = []

    if is_exact_dup:
        triggered_colors.append("Grey")
        hard_issues.append(f"B-07 - exact duplicate of row {dup_of_row} (same product, manufacturer, variant and size) - remove later row")

    dist_usa = str(row.get("Distributed in USA", "")).strip().lower()
    if dist_usa in ["no", "n", "false", "0"]:
        triggered_colors.append("Orange")
        hard_issues.append("G-08 - Distributed in USA = No - no verified US distributor")

    raw_rating = row.get("Product rating (Amazon)", None)
    if pd.notna(raw_rating):
        m_rat = re.search(r"(\d+(?:\.\d+)?)", str(raw_rating))
        if m_rat:
            r_val = float(m_rat.group(1))
            if r_val < min_rating:
                triggered_colors.append("Blue")
                r_disp = int(r_val) if r_val.is_integer() else r_val
                hard_issues.append(f"G-09 - Amazon rating {r_disp} below {min_rating} minimum - premium-fit fail")

    for pat in G07_CLAIM_PATTERNS:
        if re.search(pat, desc, re.I):
            triggered_colors.append("Pink")
            hard_issues.append("G-07 - unsubstantiated disease/miracle claim in description")
            break

    is_missing_ing = False
    matched_placeholder = None
    for pat in G01_PLACEHOLDER_PATTERNS:
        m_ph = re.search(pat, ing_l, re.I)
        if m_ph:
            is_missing_ing = True
            matched_placeholder = m_ph.group(0).strip()
            break

    if is_missing_ing:
        triggered_colors.append("Yellow")
        if is_hardware:
            hard_issues.append("G-01 - materials list missing (n/a)")
        else:
            note = f"('{matched_placeholder}')" if matched_placeholder and matched_placeholder not in ["", "nan", "n/a", "none"] else "(blank)"
            hard_issues.append(f"G-01 - full ingredient list not on file {note}")

    if not mfg_link or mfg_link.lower() in ["nan", "none", "n/a"]:
        triggered_colors.append("Yellow")
        hard_issues.append("PENDING DATA - B-06/G-08 - manufacturer link missing - complete record")

    if ing and re.search(r"\bproprietary(?:\s+blend)?\b|\bsecret blend\b", ing_l):
        triggered_colors.append("Purple")
        hard_issues.append("G-03 - contains 'proprietary' - undisclosed formula")
    elif (not is_ingestible) and re.search(r"\b(?:fragrance|parfum)\b", ing_l) and not re.search(r"phthalate[- ]free|essential oil|natural fragrance", ing_l):
        triggered_colors.append("Purple")
        hard_issues.append("G-03/P-02 - contains undisclosed 'fragrance/parfum' without phthalate-free/allergen disclosure")

    if ing and not is_hardware:
        for rule_id, pat, desc_label in BANNED_PATTERNS:
            if rule_id == "F-21" and not is_ingestible:
                continue
            m_ban = re.search(pat, ing_l)
            if m_ban:
                found_tok = m_ban.group(0)
                triggered_colors.append("Red")
                hard_issues.append(f"{rule_id} - contains {found_tok} - {desc_label}")

        if re.search(r"\b(?:sodium|potassium) benzoate\b|\be211\b|\be212\b", ing_l) and re.search(r"\b(?:ascorbic acid|citric acid|vitamin c)\b", ing_l):
            triggered_colors.append("Red")
            hard_issues.append("F-10 - benzoate preservative combined with ascorbic/citric acid (benzene formation risk)")

        if re.search(r"\b(?:bitter orange|synephrine|citrus aurantium)\b", ing_l) and re.search(r"\b(?:caffeine|guarana|green tea|yerba mate)\b", ing_l):
            triggered_colors.append("Red")
            hard_issues.append("F-19 - bitter orange / synephrine with caffeine - synephrine stack")

        if is_ingestible:
            f14_found = []
            has_explicit_artificial_flavor = False
            for m_name, m_pat in F14_MARKERS:
                if re.search(m_pat, ing_l):
                    f14_found.append(m_name)
                    if m_name == "artificial flavor":
                        has_explicit_artificial_flavor = True

            if len(f14_found) >= 2 or has_explicit_artificial_flavor:
                triggered_colors.append("Red")
                markers_str = ", ".join(f14_found)
                hard_issues.append(f"F-14 - {len(f14_found)} ultra-processed marker(s) ({markers_str}) - NOVA-4 formulation / artificial flavor ban")
            elif len(f14_found) == 1:
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - F-14 - 1 ultra-processed marker ({f14_found[0]}) - reformulation or justification needed")

    if is_ingestible and ing:
        for s_pat in W01_SUGARS:
            m_sug = re.search(s_pat, ing_l)
            if m_sug:
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - W-01 - added sugar source ({m_sug.group(0)}) - confirm <=8 g added sugar per serving (nutrition panel)")
                break

    if is_ingestible and ing:
        for c_pat in W02_CAFFEINE:
            m_caf = re.search(c_pat, ing_l)
            if m_caf:
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - W-02 - added caffeine source ({m_caf.group(0)}) - confirm <=200 mg per serving")
                break

    if is_ingestible and ing:
        for g_pat in W03_GUMS:
            m_gum = re.search(g_pat, ing_l)
            if m_gum:
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - W-03 - contains {m_gum.group(0)} - gum/thickener justification")
                break

    if is_ingestible and ing:
        m_nf = re.search(r"\b(?:natural\s+flavo(?:u)?r(?:ing)?s?|flavo(?:u)?r(?:ing)?s?)\b", ing_l)
        if m_nf and not re.search(r"artificial\s+(?:\w+\s+)?flavo", ing_l):
            triggered_colors.append("Amber")
            hold_issues.append(f"HOLD - W-04 - '{m_nf.group(0)}' without source - request source letter")

    if any(k in cat_l for k in ["grocery", "snack", "beverage", "coffee", "tea", "food", "pantry"]):
        w06_crops = detect_w06_conventional(ing, prod_name=p_name)
        for crop in w06_crops[:2]:
            triggered_colors.append("Amber")
            hold_issues.append(f"HOLD - W-06 - conventional {crop} (no organic claim) - prefer USDA Organic or residue test")

    if is_ingestible:
        combined_search = f"{p_name} {ing}".lower()
        for trig_label, trig_pat in G06_HEAVY_METAL_TRIGGERS:
            if re.search(trig_pat, combined_search):
                triggered_colors.append("Amber")
                hm_type = "heavy metals" if trig_label in ["matcha", "tea", "cardamom"] else "heavy metals + mycotoxins"
                hold_issues.append(f"HOLD - G-06/W-07/F-16 - COA required ({hm_type}) - trigger: {trig_label}")
                break

    if is_ingestible and ing:
        for auth_item in F17_AUTHENTICITY_TRIGGERS:
            if re.search(r"\b" + re.escape(auth_item) + r"\b", ing_l):
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - F-17 - {auth_item} - authenticity + contaminant proof (COA) required")
                break

    if any(k in cat_l for k in ["coffee", "tea", "beverage"]):
        if re.search(r"\b(?:decaf|decaffeinated|half-caff)\b", f"{p_name} {ing_l}") and not re.search(r"swiss water|co2|ethyl acetate|sugarcane", f"{p_name} {ing_l} {desc.lower()}"):
            triggered_colors.append("Amber")
            hold_issues.append("HOLD - F-15 - decaffeinated - decaffeination method proof (Swiss Water / CO2 / ethyl acetate) required")

    if re.search(r"\b(?:dark roast|french roast|potato chips|crispbread)\b", f"{p_name} {desc}".lower()):
        triggered_colors.append("Amber")
        hold_issues.append("HOLD - W-08 - dark roast / fried item - request acrylamide test")

    if "coffee" in cat_l or "tea" in cat_l:
        m_bag = re.search(r"\b(pyramid|teabag|tea bag|pod|capsule|k-cup|sachet)\b", f"{p_name} {sizes} {desc}".lower())
        if m_bag:
            triggered_colors.append("Amber")
            hold_issues.append(f"HOLD - W-09 - {m_bag.group(1)} - confirm plastic-free (paper/muslin/PLA) bag or pod material")

    if is_ingestible:
        m_alc = re.search(r"\b(kombucha|kefir|beer|wine|alcohol|fermented)\b", f"{p_name} {ing_l}".lower())
        if m_alc:
            triggered_colors.append("Amber")
            hold_issues.append(f"HOLD - W-10 - {m_alc.group(1)} - confirm <=0.5% ABV for non-alcoholic labeling")

    if is_ingestible and re.search(r"\b(?:silicon dioxide|silica|magnesium stearate)\b", ing_l):
        triggered_colors.append("Amber")
        hold_issues.append("HOLD - W-11/F-20 - contains silicon dioxide / stearate excipient - disclose source / level")

    if is_ingestible and ing:
        for v_name, v_pat in W12_UPPER_LIMIT_NUTRIENTS:
            if re.search(v_pat, ing_l):
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - W-12 - {v_name} - confirm dose below tolerable upper limit")
                break

    if is_oral and re.search(r"\b(?:sodium fluoride|stannous fluoride|fluoride)\b", ing_l):
        triggered_colors.append("Amber")
        hold_issues.append("HOLD - W-14 - contains fluoride - assortment / hydroxyapatite pairing decision")

    m_can = re.search(r"\b(\d+(?:\.\d+)?\s*(?:oz|g|ml)?\s*can|canned|in a can|tin)\b", f"{p_name} {sizes}".lower())
    if m_can and is_ingestible:
        triggered_colors.append("Amber")
        hold_issues.append(f"HOLD - W-17/M-01 - can/tin packaging ({m_can.group(1)}) - BPA-NI can-lining proof required")

    if is_ingestible and ing:
        if re.search(r"\bturmeric\b", ing_l) and re.search(r"\b(?:black pepper|piperine)\b", ing_l):
            triggered_colors.append("Amber")
            hold_issues.append("HOLD - W-18 - turmeric with piperine/black pepper - add caution label")
        for b_name, b_pat in W18_BOTANICAL_CAUTIONS:
            if re.search(b_pat, ing_l):
                triggered_colors.append("Amber")
                hold_issues.append(f"HOLD - W-18 - contains {b_name} - liver/hormonal caution label required")

    if sibling_rows and not is_exact_dup:
        triggered_colors.append("Amber")
        rows_str = ", ".join(str(r) for r in sibling_rows[:6])
        hold_issues.append(f"HOLD - B-07 - same product family as rows {rows_str} - group variants into one listing")

    if is_supplement:
        triggered_colors.append("Amber")
        hold_issues.append("HOLD - C-03/G-06 - third-party certification (NSF / USP / Informed Choice / ConsumerLab), cGMP proof and batch COA required")
    elif "wearables" in cat_l:
        triggered_colors.append("Amber")
        hold_issues.append("HOLD - M-05/M-06/M-07/M-03/M-04 - FCC/UL proof, FDA clearance for any medical claims, privacy policy and band-material (PFAS/nickel) declarations required")

    if not triggered_colors:
        return "Passed", "None", "Passed - no flags"

    winning_color = min(triggered_colors, key=lambda c: COLOR_HIERARCHY[c]["priority"])
    win_meta = COLOR_HIERARCHY[winning_color]
    final_status = win_meta["status"]
    base_prefix = win_meta["prefix"]

    all_issues = hard_issues + hold_issues
    if winning_color == "Amber" and all_issues:
        first_item = re.sub(r"^HOLD\s*[-]\s*", "", all_issues[0])
        formatted_body = "; ".join([first_item] + all_issues[1:])
        full_summary = f"{base_prefix}: {formatted_body}"
    else:
        extra_hard = len(hard_issues) - 1
        if winning_color in ["Red", "Purple"] and extra_hard > 0:
            prefix_str = f"{base_prefix} (+{extra_hard} more)"
        else:
            prefix_str = base_prefix
        formatted_body = "; ".join(all_issues)
        full_summary = f"{prefix_str}: {formatted_body}"

    return final_status, winning_color, full_summary


def tool_screen_and_writeback(wb_source, min_rating: float = 4.2, prefer_existing_notes: bool = False):
    wb = openpyxl.load_workbook(wb_source)
    all_screened_records = []
    category_dfs = {}
    sheet_sku_counts = {}

    for idx, sheet_name in enumerate(wb.sheetnames):
        if sheet_name == "NN Shortlisting Guidelines":
            continue

        ws = wb[sheet_name]

        header_row_idx = None
        for r_scan in range(1, min(26, ws.max_row + 1)):
            row_vals = [
                str(ws.cell(row=r_scan, column=c).value or "").strip().lower()
                for c in range(1, min(15, ws.max_column + 1))
            ]
            if "product name" in row_vals:
                header_row_idx = r_scan
                break

        if not header_row_idx:
            sheet_sku_counts[sheet_name] = 0
            category_dfs[sheet_name] = pd.DataFrame()
            continue

        data_start_row = header_row_idx + 1

        col_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=header_row_idx, column=c).value
            if val is not None:
                col_map[str(val).strip()] = c

        prod_col = col_map.get("Product name")
        if not prod_col:
            sheet_sku_counts[sheet_name] = 0
            category_dfs[sheet_name] = pd.DataFrame()
            continue

        status_col = col_map.get("Screening issue") or col_map.get("Screening Status")
        if not status_col:
            status_col = ws.max_column + 1
            ws.cell(row=header_row_idx, column=status_col).value = "Screening issue"
            ws.cell(row=header_row_idx, column=status_col).font = Font(bold=True)

        raw_rows = []
        exact_seen = {}
        family_map = {}

        for r in range(data_start_row, ws.max_row + 1):
            prod_val = ws.cell(row=r, column=prod_col).value
            if prod_val is None:
                continue
            prod_str = str(prod_val).strip()
            if not prod_str or re.match(r"^(?:As per manufacturer|-{1,2}\s*-{1,2})$", prod_str, re.I):
                continue

            row_data = {"Category": sheet_name, "Product name": prod_str, "_excel_row": r}
            for col_name, c_idx in col_map.items():
                row_data[col_name] = ws.cell(row=r, column=c_idx).value

            mfg = str(row_data.get("Manufacturer name", "")).strip().lower()
            var = str(row_data.get("Flavor / Variant", "")).strip().lower()
            sz = str(row_data.get("Sizes available", "")).strip().lower()
            p_low = prod_str.lower()

            exact_key = (p_low, mfg, var, sz)
            fam_key = (p_low, mfg)

            if exact_key in exact_seen:
                row_data["_is_exact_dup"] = True
                row_data["_dup_of_row"] = exact_seen[exact_key]
            else:
                exact_seen[exact_key] = r
                row_data["_is_exact_dup"] = False
                row_data["_dup_of_row"] = None

            family_map.setdefault(fam_key, []).append(r)
            raw_rows.append(row_data)

        sheet_records = []
        for row_data in raw_rows:
            r = row_data["_excel_row"]
            p_low = row_data["Product name"].lower()
            mfg = str(row_data.get("Manufacturer name", "")).strip().lower()
            siblings = [x for x in family_map.get((p_low, mfg), []) if x != r]

            status, p_color, full_summary = evaluate_product_comprehensive(
                row_data,
                min_rating=min_rating,
                is_exact_dup=row_data["_is_exact_dup"],
                dup_of_row=row_data["_dup_of_row"],
                sibling_rows=siblings,
                prefer_existing_notes=prefer_existing_notes
            )

            row_data["Screening Status"] = status
            row_data["Priority Color"] = p_color
            row_data["Screening Summary"] = full_summary
            sheet_records.append(row_data)

            target_cell = ws.cell(row=r, column=status_col)
            target_cell.value = full_summary
            c_meta = COLOR_HIERARCHY.get(p_color, COLOR_HIERARCHY["None"])
            target_cell.fill = PatternFill(start_color=c_meta["fill"], end_color=c_meta["fill"], fill_type="solid")
            target_cell.font = Font(color=c_meta["font"], bold=True)
            target_cell.alignment = Alignment(wrap_text=True, vertical="center")

        sheet_sku_counts[sheet_name] = len(sheet_records)
        df_sheet = pd.DataFrame(sheet_records) if sheet_records else pd.DataFrame()
        category_dfs[sheet_name] = df_sheet
        if not df_sheet.empty:
            all_screened_records.append(df_sheet)

    master_df = pd.concat(all_screened_records, ignore_index=True) if all_screened_records else pd.DataFrame()

    if not master_df.empty:
        p_col = "Price/unit" if "Price/unit" in master_df.columns else "Price (MRP)-Manuf"
        if p_col in master_df.columns:
            master_df["Numeric_Price"] = master_df[p_col].astype(str).str.extract(r"(\d+(?:\.\d+)?)")[0].astype(float)
        else:
            master_df["Numeric_Price"] = None

        if "MOQ" in master_df.columns:
            master_df["Numeric_MOQ"] = master_df["MOQ"].astype(str).str.extract(r"(\d+)")[0].astype(float)
        else:
            master_df["Numeric_MOQ"] = None

    out_stream = BytesIO()
    wb.save(out_stream)
    out_stream.seek(0)

    return master_df, category_dfs, sheet_sku_counts, out_stream


def tool_lookup_product(df, query):
    stop_words = {
        "why", "was", "is", "the", "product", "marked", "as", "in", "for",
        "check", "status", "of", "tell", "me", "about", "audit", "screen",
        "hold", "rejected", "banned", "passed", "draft", "email", "outreach"
    }
    tokens = [t for t in re.split(r"\W+", query.lower()) if len(t) > 2 and t not in stop_words]
    if not tokens:
        return pd.DataFrame()

    def score_name(name):
        n_low = str(name).lower()
        return sum(1 for tok in tokens if tok in n_low)

    scores = df["Product name"].apply(score_name)
    matched = df[scores > 0].copy()
    if matched.empty:
        return pd.DataFrame()
    matched["_score"] = scores[scores > 0]
    return matched.sort_values(by="_score", ascending=False).drop(columns=["_score"]).head(5)


def tool_category_summary(df, category_name):
    cat_df = df[df["Category"].str.lower() == category_name.lower()]
    if cat_df.empty:
        return None
    tot = len(cat_df)
    passed = int((cat_df["Screening Status"] == "Passed").sum())
    return {
        "Category": category_name,
        "Total": tot,
        "Passed": passed,
        "Pass_Rate": (passed / tot) * 100 if tot else 0,
        "Hold": int((cat_df["Screening Status"] == "Hold (Proof Required)").sum()),
        "Rejected": int((cat_df["Screening Status"] == "Rejected").sum()),
        "Pending Data": int((cat_df["Screening Status"] == "Pending Data").sum())
    }


def tool_chemical_toxin_summary(df, top_n=12):
    rej_df = df[df["Screening Status"] == "Rejected"]
    if rej_df.empty:
        return pd.DataFrame(columns=["Screening Rule / Toxin", "Violations Count"])
    all_rules = []
    for summary in rej_df["Screening Summary"].dropna():
        body = str(summary).split(":", 1)[-1]
        for item in body.split(";"):
            item_clean = item.strip()
            if any(item_clean.startswith(p) for p in ["F-", "P-", "H-", "M-", "G-03", "G-07", "G-08", "G-09", "B-07"]):
                all_rules.append(item_clean)
    res = pd.Series(all_rules).value_counts().head(top_n).reset_index()
    res.columns = ["Screening Rule / Toxin", "Violations Count"]
    return res


def tool_generate_coa_email(prod_series):
    m_name = prod_series.get("Manufacturer name") or "Vendor Partner"
    p_name = prod_series.get("Product name") or "Product"
    issues = prod_series.get("Screening Summary") or "COA Documentation"
    email_addr = prod_series.get("Contact Email") or "compliance@brand.com"

    body = (
        f"Subject: Nutrition Nest Onboarding: Batch COA & Documentation Request for {p_name}\n\n"
        f"Dear {m_name} Quality & Regulatory Team,\n\n"
        f"We are currently evaluating {p_name} for procurement and retail shortlisting at Nutrition Nest (www.nutritionnest.com).\n\n"
        f"To clear our Stage-2 safety gates under our strict ingredient standards, we kindly request the following documentation for this SKU:\n"
        f"- Batch-level Certificate of Analysis (COA) & verification for: {issues}\n"
        f"- Lab testing verification for heavy metals (lead, arsenic, cadmium, mercury), mycotoxins, and pesticide/glyphosate residues.\n"
        f"- Confirmation of FDA facility registration and product liability insurance.\n\n"
        f"Please forward the testing documentation at your earliest convenience so our regulatory committee can finalize product selection.\n\n"
        f"Warm regards,\n\n"
        f"Category Procurement Team\n"
        f"Nutrition Nest"
    )
    return email_addr, body


def tool_commercial_metrics(df):
    clean = df[df["Screening Status"] == "Passed"]
    if clean.empty:
        return {"Total": 0, "Median_Price": 0.0, "Median_MOQ": 1}
    p_med = clean["Numeric_Price"].dropna().median() if "Numeric_Price" in clean.columns and not clean["Numeric_Price"].dropna().empty else 0.0
    m_med = clean["Numeric_MOQ"].dropna().median() if "Numeric_MOQ" in clean.columns and not clean["Numeric_MOQ"].dropna().empty else 1.0
    return {
        "Total": len(clean),
        "Median_Price": float(p_med),
        "Median_MOQ": int(m_med)
    }
